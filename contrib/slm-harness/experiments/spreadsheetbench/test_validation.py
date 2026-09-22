import tempfile,unittest,zipfile
from pathlib import Path
import openpyxl
from workbook_validation import validate_output
class ValidationTests(unittest.TestCase):
 def setUp(self):
  self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name);self.work=self.root/'work';self.work.mkdir()
  self.input=self.root/'input.xlsx';w=openpyxl.Workbook();w.save(self.input)
 def tearDown(self):self.temp.cleanup()
 def test_expanded_archive_is_rejected_before_parsing(self):
  with zipfile.ZipFile(self.work/'output.xlsx','w',zipfile.ZIP_DEFLATED) as z:
   z.writestr('xl/workbook.xml','x'*(17*1024*1024))
  self.assertLess((self.work/'output.xlsx').stat().st_size,100000)
  result=validate_output(self.input,self.work);self.assertFalse(result['ok'])
  self.assertIn('archive expansion limit',result['stderr'])
 def test_valid_workbook_is_accepted_but_fake_structure_is_not(self):
  w=openpyxl.Workbook();w.active['A1']='hello';w.save(self.work/'output.xlsx')
  self.assertTrue(validate_output(self.input,self.work)['ok'])
  with zipfile.ZipFile(self.work/'output.xlsx','w') as z:z.writestr('xl/workbook.xml','not a workbook')
  self.assertFalse(validate_output(self.input,self.work)['ok'])
if __name__=='__main__':unittest.main()

class WorkerContractTests(unittest.TestCase):
 def test_failed_submit_does_not_announce_registration(self):
  import json,subprocess,sys
  with tempfile.TemporaryDirectory() as tmp:
   root=Path(tmp);openpyxl.Workbook().save(root/'input.xlsx')
   r=subprocess.run([sys.executable,str(Path(__file__).with_name('adapter_v2.py')),tmp,'submit_solution'],input=json.dumps({'code':"print('no output')"}),text=True,capture_output=True,check=True)
   result=json.loads(r.stdout)
   self.assertFalse(result['validated'])
   self.assertNotIn('Submission registered',result['note'])
 def test_valid_submit_runs_bounded_validation_and_registers(self):
  import json,subprocess,sys
  with tempfile.TemporaryDirectory() as tmp:
   root=Path(tmp);openpyxl.Workbook().save(root/'input.xlsx')
   code="import openpyxl;w=openpyxl.load_workbook('/input/input.xlsx');w.active['A1']=42;w.save('/work/output.xlsx')"
   r=subprocess.run([sys.executable,str(Path(__file__).with_name('adapter_v2.py')),tmp,'submit_solution'],input=json.dumps({'code':code}),text=True,capture_output=True,check=True)
   result=json.loads(r.stdout)
   self.assertTrue(result['validated']);self.assertTrue(result['validation']['ok'])
   self.assertEqual((root/'solution.py').read_text(),code)
