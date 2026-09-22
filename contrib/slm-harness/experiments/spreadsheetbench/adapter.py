"""Trusted benchmark adapter. Gold is used only by compare_output, never by tools."""
import contextlib
import io
import json
import os
from pathlib import Path
import re
import resource
import shutil
import signal
import subprocess
import sys
import tempfile
import zipfile

HERE=Path(__file__).resolve().parent
ROOT=HERE.parents[1]
RUNTIME=Path('/home/hugo/.cache/codex-runtimes/codex-primary-runtime/dependencies/python')
UPSTREAM=ROOT/'artifacts/spreadsheetbench/upstream'

def execute_python(input_path, work, code, fresh=False, timeout=25, read_only=False):
    work=Path(work).resolve();input_path=Path(input_path).resolve()
    if fresh and work.exists():shutil.rmtree(work)
    work.mkdir(parents=True,exist_ok=True)
    control=tempfile.TemporaryDirectory(prefix='sb-trusted-')
    script=Path(control.name)/'execute.py';script.write_text(code)
    cmd=['/usr/bin/bwrap','--die-with-parent','--new-session','--unshare-all','--clearenv',
         '--ro-bind','/usr','/usr','--ro-bind','/lib','/lib','--ro-bind','/lib64','/lib64',
         '--ro-bind',str(RUNTIME),str(RUNTIME),'--proc','/proc','--dev','/dev','--tmpfs','/tmp',
         '--dir','/input','--ro-bind',str(input_path),'/input/input.xlsx',
         '--ro-bind' if read_only else '--bind',str(work),'/work','--ro-bind',str(script),'/task_code.py','--chdir','/work',
         '--setenv','HOME','/tmp','--setenv','PATH',f'{RUNTIME}/bin:/usr/bin',
         '--setenv','OPENBLAS_NUM_THREADS','1','--setenv','OMP_NUM_THREADS','1',
         str(RUNTIME/'bin/python3'),'-I','/task_code.py']
    def limits():
        resource.setrlimit(resource.RLIMIT_AS,(3*1024**3,3*1024**3))
        resource.setrlimit(resource.RLIMIT_CPU,(20,20))
        resource.setrlimit(resource.RLIMIT_FSIZE,(32*1024**2,32*1024**2))
        resource.setrlimit(resource.RLIMIT_NOFILE,(128,128))
        resource.setrlimit(resource.RLIMIT_CORE,(0,0))
    with tempfile.TemporaryFile() as out,tempfile.TemporaryFile() as err:
        try:
            p=subprocess.Popen(cmd,stdout=out,stderr=err,preexec_fn=limits,start_new_session=True)
            try:p.wait(timeout=timeout);exit_code=p.returncode
            except subprocess.TimeoutExpired:
                os.killpg(p.pid,signal.SIGKILL);p.wait();exit_code=124
            out.seek(0);err.seek(0)
            stdout=out.read(12000).decode(errors='replace');stderr=err.read(12000).decode(errors='replace')
        finally:control.cleanup()
    output=work/'output.xlsx'
    valid=output.is_file() and not output.is_symlink() and output.stat().st_size<=32*1024**2
    if valid:
        try:
            with zipfile.ZipFile(output) as z:valid='xl/workbook.xml' in z.namelist()
        except Exception:valid=False
    return {'exit_code':exit_code,'stdout':stdout,'stderr':stderr,'output_valid':valid}

def input_gold(folder):
    folder=Path(folder)
    inputs=sorted(folder.glob('*_init.xlsx')) or sorted(folder.glob('initial.xlsx'))
    gold=sorted(folder.glob('*_golden.xlsx')) or sorted(folder.glob('golden.xlsx'))
    if len(inputs)!=1 or len(gold)!=1:raise ValueError(f'Expected one verified pair: {folder.name}')
    return inputs[0],gold[0]

def resolve_regions(task,gold,pred):
    from openpyxl.utils.cell import range_boundaries,get_column_letter
    result=[]
    position=task['answer_position']
    # Verified metadata typo, supported by instruction's column BD and data_position.
    if str(task.get('id'))=='73-45' and position=="'Sheet1'!BD2:308":
        position="'Sheet1'!BD2:BD308"
    pattern=re.compile(r"(?:^|,)\s*(?:(.*?)!)?'?(\$?[A-Z]+\$?\d*(?::\$?[A-Z]*\$?\d*)?)'?\s*(?=,|$)",re.I)
    matches=list(pattern.finditer(position))
    if not matches or ''.join(m.group(0) for m in matches)!=position:
        raise ValueError(f'Unsupported answer_position: {position}')
    for match in matches:
        part=(match.group(1)+'!' if match.group(1) is not None else '')+match.group(2)
        if '!' in part:sheet,region=part.rsplit('!',1);sheet=sheet.strip().strip("'")
        else:
            sheet=task.get('answer_sheet','').split(',')[0].strip().strip("'") or gold.sheetnames[0]
            region=part
        region=region.strip().strip("'").replace('$','')
        bounds=range_boundaries(region);a,b,c,d=bounds
        if sheet not in gold:raise ValueError(f'Unknown gold sheet {sheet}')
        ws=gold[sheet];ps=pred[sheet] if sheet in pred else ws
        a=a or 1;b=b or 1;c=c or max(ws.max_column,ps.max_column);d=d or max(ws.max_row,ps.max_row)
        if (c-a+1)*(d-b+1)>1000000:raise ValueError('evaluation range exceeds 1M cells')
        result.append((sheet,f'{get_column_letter(a)}{b}:{get_column_letter(c)}{d}'))
    return result

def compare_output(gold_path,output_path,task):
    import openpyxl
    if not Path(output_path).is_file():return {'passed':False,'reason':'missing_output'}
    sys.path.insert(0,str(UPSTREAM/'evaluation'))
    from evaluation import cell_level_compare
    gold=pred=None
    try:
        gold=openpyxl.load_workbook(gold_path,data_only=True);pred=openpyxl.load_workbook(output_path,data_only=True)
        regions=resolve_regions(task,gold,pred);checks=[]
        for sheet,region in regions:
            with contextlib.redirect_stdout(io.StringIO()):ok,msg=cell_level_compare(gold,pred,sheet,region)
            checks.append({'sheet':sheet,'range':region,'passed':ok,'detail':msg})
        return {'passed':all(c['passed'] for c in checks),'checks':checks}
    except Exception as exc:return {'passed':False,'reason':'grader_error','error':repr(exc)}
    finally:
        if gold:gold.close()
        if pred:pred.close()

if __name__=='__main__':
    # CLI bridge is launched by the generic jingwei host with system Python.
    # Re-exec trusted adapter in its local dependency environment.
    expected=HERE/'.venv/bin/python'
    if Path(sys.prefix).resolve()!=(HERE/'.venv').resolve():os.execv(str(expected),[str(expected),__file__,*sys.argv[1:]])
    task_dir=Path(sys.argv[1]).resolve();tool=sys.argv[2];args=json.load(sys.stdin)
    code=args.get('code')
    if not isinstance(code,str) or len(code)>60000:raise ValueError('code must be <=60000 characters')
    if tool not in ('inspect_python','submit_solution'):raise ValueError('unknown tool')
    submit=tool=='submit_solution'
    result=execute_python(task_dir/'input.xlsx',task_dir/'sandbox',code,fresh=submit)
    if submit:
        (task_dir/'solution.py').write_text(code)
        result['ok']=result['exit_code']==0 and result['output_valid']
        result['validated']=result['ok']
        result['note']='File created; correctness is evaluated separately. Submit corrected standalone code if necessary.'
    print(json.dumps(result,ensure_ascii=False))
