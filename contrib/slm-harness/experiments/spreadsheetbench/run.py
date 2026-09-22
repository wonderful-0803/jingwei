"""Resumable jingwei + Ollama evaluation on all Verified tasks, with offline grading."""
import argparse
import collections
import fcntl
import hashlib
import json
import os
from pathlib import Path
import selectors
import shutil
import signal
import subprocess
import sys
import time
import urllib.request
import uuid
import openpyxl
import numpy
import pandas
from adapter import HERE,ROOT,input_gold,execute_python,compare_output
from proxy import proxy
from sync_state import record
from make_report import make_report

DATA=ROOT/'artifacts/spreadsheetbench/spreadsheetbench_verified_400'
BINARY=ROOT/'jingwei/target/debug/jingwei-bench-runner'
MODELS={'9b':'hf.co/unsloth/Qwen3.5-9B-GGUF:Q8_0','27b':'hf.co/unsloth/Qwen3.8-27B-GGUF:UD-IQ3_S'}

def dump(path,value):
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
    tmp=path.with_suffix(path.suffix+'.tmp');tmp.write_text(json.dumps(value,ensure_ascii=False,indent=2,default=str));tmp.replace(path)

def api(path,payload=None):
    req=urllib.request.Request('http://127.0.0.1:11434/api/'+path,None if payload is None else json.dumps(payload).encode(),{'Content-Type':'application/json'})
    with urllib.request.urlopen(req,timeout=600) as response:return json.load(response)

def sha(path):return hashlib.sha256(Path(path).read_bytes()).hexdigest()

def read_jsonl(path):
    # str.splitlines() also splits U+0085, which is legal inside a JSON string.
    return [json.loads(line) for line in Path(path).read_text().split('\n') if line]

class Calc:
    def __init__(self,root):
        self.root=Path(root);self.root.mkdir(parents=True,exist_ok=True)
        self.log=(self.root/'worker.log').open('ab')
        self.proc=subprocess.Popen(['/opt/libreoffice26.8/program/python',str(HERE/'recalculate.py'),str(self.root)],stdin=subprocess.PIPE,stdout=subprocess.PIPE,stderr=self.log,text=True,bufsize=1,start_new_session=True)
        try:
            ready=self.read()
            if not ready.get('ready'):raise RuntimeError(ready)
        except BaseException:
            self.close();raise
    def read(self):
        with selectors.DefaultSelector() as selector:
            selector.register(self.proc.stdout,selectors.EVENT_READ)
            if not selector.select(90):raise TimeoutError('LibreOffice worker timeout')
            line=self.proc.stdout.readline()
            if not line:raise RuntimeError('LibreOffice worker exited')
            return json.loads(line)
    def run(self,source,dest):
        self.proc.stdin.write(json.dumps({'input':str(source),'output':str(dest)})+'\n');self.proc.stdin.flush();return self.read()
    def close(self):
        if self.proc.poll() is None:
            self.proc.stdin.close()
            try:self.proc.wait(timeout=10)
            except subprocess.TimeoutExpired:os.killpg(self.proc.pid,signal.SIGKILL);self.proc.wait()
        self.log.close()

def preview(path):
    book=openpyxl.load_workbook(path,read_only=True,data_only=False);out=[]
    try:
        for sheet in book:
            max_row=sheet.max_row or 1
            max_column=sheet.max_column or 1
            out.append({'sheet':sheet.title,'max_row':sheet.max_row,'max_column':sheet.max_column,'first_five_rows':list(sheet.iter_rows(min_row=1,max_row=min(5,max_row),max_col=min(40,max_column),values_only=True))})
    finally:book.close()
    return json.dumps(out,ensure_ascii=False,default=str)[:16000]

def prompt(task,path):
    return ('Solve this SpreadsheetBench task. Save the solution workbook to /work/output.xlsx.\n'
            'Input workbook: /input/input.xlsx\nInstruction:\n'+task['instruction']+
            '\nInstruction type: '+task['instruction_type']+'\nAnswer position: '+task['answer_position']+
            '\nAnswer sheet metadata (if provided): '+task.get('answer_sheet','')+
            '\nWorkbook preview (first five rows per sheet, at most 40 columns; inspect more as needed):\n'+preview(path))

def prepare_attempt(folder):
    if folder.exists() and any(folder.iterdir()):
        archive=folder.parents[1]/'.interrupted'/folder.parent.name/(folder.name+'-'+uuid.uuid4().hex)
        archive.parent.mkdir(parents=True,exist_ok=True);folder.rename(archive)
    folder.mkdir(parents=True,exist_ok=True)


def task_run(task,folder,model,calc,profile=None):
    prepare_attempt(folder)
    folder.mkdir(parents=True,exist_ok=True);inp,gold=input_gold(DATA/task['spreadsheet_path'])
    shutil.copyfile(inp,folder/'input.xlsx')
    started=time.monotonic()
    with proxy(folder/'llm-trace.jsonl',merge_system=True) as endpoint:
        config={'profile':str(profile or HERE/'profile.json'),'task_dir':str(folder),'worker_args':[str(folder)],'endpoint':endpoint,'model':model,'prompt':prompt(task,inp)}
        dump(folder/'config.json',config)
        with (folder/'host.log').open('wb') as log:
            p=subprocess.Popen([str(BINARY),str(folder/'config.json')],stdout=log,stderr=subprocess.STDOUT,start_new_session=True)
            try:p.wait(timeout=660);host_exit=p.returncode
            except subprocess.TimeoutExpired:os.killpg(p.pid,signal.SIGKILL);p.wait();host_exit=124
    result={'id':str(task['id']),'instruction_type':task['instruction_type'],'model':model,'host_exit':host_exit,'passed':False}
    report_path=folder/'agent-report.json'
    if report_path.exists():result['agent']=json.loads(report_path.read_text())
    solution=folder/'solution.py'
    if solution.exists():
        execution=execute_python(inp,folder/'replay',solution.read_text(),fresh=True)
        dump(folder/'execution.json',execution);result['execution']=execution
        if execution['exit_code']==0 and execution['output_valid']:
            recalc=calc.run(folder/'replay/output.xlsx',folder/'recalculated.xlsx');result['recalculation']=recalc
            if recalc.get('ok'):
                grade=compare_output(gold,folder/'recalculated.xlsx',task);result['grade']=grade;result['passed']=grade['passed']
                result['category']='passed' if grade['passed'] else grade.get('reason','wrong_answer')
            else:result['category']='recalculation_error'
        else:result['category']='execution_error' if execution['exit_code'] else 'missing_output'
    else:result['category']='no_submission'
    trace=folder/'llm-trace.jsonl'
    if trace.exists():
        calls=read_jsonl(trace);result['model_calls']=len(calls)
        result['usage']={key:sum(c.get('response',{}).get('usage',{}).get(key,0) or 0 for c in calls) for key in ('prompt_tokens','completion_tokens','total_tokens')}
        result['http_errors']=[c['status'] for c in calls if c['status']!=200]
    journal=list((folder/'journal').glob('*.jsonl'))
    result['agent_errors']=[]
    for path in journal:
        for record in read_jsonl(path):
            event=record.get('kind',{})
            if event.get('type')=='error':result['agent_errors'].append(event)
    codes=[event.get('code') for event in result['agent_errors']]
    if 'reference_result_view' in codes or 'reference_policy' in codes:result['category']='harness_error';result['passed']=False
    elif result['category']=='no_submission' and 'reference_invalid_action' in codes:result['category']='invalid_action'
    result['loaded_after_requests']=api('ps')
    result['elapsed_seconds']=time.monotonic()-started
    dump(folder/'result.json',result);return result

def summarize(root,keys,total):
    summary={'expected_per_model':total,'models':{}}
    for key in keys:
        results=[json.loads(p.read_text()) for p in (root/key).glob('*/result.json')]
        summary['models'][key]={'completed':len(results),'passed':sum(r['passed'] for r in results),'accuracy_completed':sum(r['passed'] for r in results)/len(results) if results else None,'categories':dict(collections.Counter(r['category'] for r in results)),'seconds':sum(r['elapsed_seconds'] for r in results),'model_calls':sum(r.get('model_calls',0) for r in results)}
    dump(root/'progress.json',summary);return summary

def main():
    parser=argparse.ArgumentParser();parser.add_argument('--out',required=True);parser.add_argument('--models',default='9b,27b');parser.add_argument('--ids',default='');parser.add_argument('--limit',type=int);parser.add_argument('--profile',type=Path,default=HERE/'profile.json');args=parser.parse_args()
    if args.limit is not None and (args.limit<1 or args.limit>400):parser.error('--limit must be 1..400')
    if args.limit is not None and args.ids:parser.error('choose --limit or --ids')
    root=Path(args.out).resolve();root.mkdir(parents=True,exist_ok=True);keys=args.models.split(',')
    run_lock=(root/'.runner.lock').open('a+')
    fcntl.flock(run_lock.fileno(),fcntl.LOCK_EX|fcntl.LOCK_NB)
    data=json.loads((DATA/'dataset.json').read_text())
    if args.ids:data=[t for t in data if str(t['id']) in args.ids.split(',')]
    if args.limit is not None:data=data[:args.limit]
    scope='prefix'+str(args.limit) if args.limit is not None else ('full400' if not args.ids else 'pilot')
    tags={m['name']:m for m in api('tags')['models']}
    manifest={'dataset_sha256':sha(DATA/'dataset.json'),'dataset_files_sha256':{str(p.relative_to(DATA)):sha(p) for p in DATA.rglob('*.xlsx')},'archive_sha256':sha(DATA.parent/'verified400.tar.gz'),'ids':[str(t['id']) for t in data],
              'models':{key:{'name':MODELS[key],'digest':tags[MODELS[key]]['digest']} for key in keys},
              'source_sha256':{str(p.relative_to(HERE)):sha(p) for p in HERE.rglob('*') if p.is_file() and '.venv' not in p.parts and '__pycache__' not in p.parts and p.suffix in ('.rs','.py','.md','.json','.toml','.lock')},
              'binary_sha256':sha(BINARY),'active_profile':{'path':str(args.profile.resolve()),'sha256':sha(args.profile)},'sampling':{'temperature':0,'seed':7,'reasoning_effort':'none','num_ctx':32768,'output_reserve':4096,'max_steps':8},
              'engine':subprocess.check_output(['/opt/libreoffice26.8/program/soffice','--version'],text=True).strip(),'python':sys.version,'dependencies':{'openpyxl':openpyxl.__version__,'numpy':numpy.__version__,'pandas':pandas.__version__},'upstream_commit':'49b73a94775fb489063f60ca1865e3a650079a79','mode':scope}
    old=root/'manifest.json'
    if old.exists() and json.loads(old.read_text())!=manifest:raise RuntimeError('Manifest mismatch: use a new output directory')
    dump(old,manifest);dump(root/'process.json',{'pid':os.getpid(),'started_utc':time.strftime('%Y-%m-%dT%H:%M:%SZ',time.gmtime()),'argv':sys.argv})
    if api('ps')['models']:raise RuntimeError('Ollama has another loaded model; not disturbing it')
    calc=Calc(root/'calc');summary={}
    try:
        for key in keys:
            pending=[t for t in data if not (root/key/str(t['id'])/'result.json').exists()]
            if not pending:continue
            model=MODELS[key]
            try:
                warm=api('generate',{'model':model,'prompt':'OK','stream':False,'think':False,'keep_alive':'30m','options':{'temperature':0,'seed':7,'num_ctx':32768,'num_predict':1}})
                dump(root/key/'environment.json',{'loaded':api('ps'),'version':api('version'),'gpu':subprocess.check_output(['nvidia-smi','--query-gpu=name,memory.used,memory.total','--format=csv,noheader'],text=True),'warmup':{k:v for k,v in warm.items() if k not in ('thinking','response','context')}})
                print(json.dumps({'event':'model_started','model':key,'pending':len(pending)}),flush=True)
                errors=0
                for index,task in enumerate(pending):
                    print(json.dumps({'event':'task_started','model':key,'id':str(task['id']),'index':index+1,'total':len(pending)}),flush=True)
                    execution=scope+'-'+key+'-'+str(task['id'])
                    if not args.ids:record('execution_started',execution,'运行 '+key+' 题 '+str(task['id'])+'；含真实模型/工具调用、独立重放、重算、盲评',str((root/key/str(task['id'])).relative_to(ROOT)))
                    result=task_run(task,root/key/str(task['id']),model,calc,profile=args.profile.resolve())
                    if not args.ids:record('execution_finished',execution,'完成 '+key+' 题 '+str(task['id'])+'；结果 '+result['category']+'，耗时 '+str(round(result['elapsed_seconds'],1))+'s',str((root/key/str(task['id'])/'result.json').relative_to(ROOT)),result['host_exit'])
                    errors=errors+1 if result.get('http_errors') else 0
                    summary=summarize(root,keys,len(data))
                    print(json.dumps({'event':'task_finished','model':key,'id':result['id'],'passed':result['passed'],'category':result['category'],'seconds':round(result['elapsed_seconds'],1),'progress':summary['models'][key]},ensure_ascii=False),flush=True)
                    if result['category']=='harness_error':raise RuntimeError('Harness infrastructure error; stop and diagnose')
                    if errors>=3:raise RuntimeError('Three consecutive tasks with transport failures; stopping rather than scoring outage')
            finally:
                api('generate',{'model':model,'stream':False,'keep_alive':0})
                dump(root/key/'cleanup.json',{'loaded':api('ps')})
        summary=summarize(root,keys,len(data));dump(root/'complete.json',summary)
        if not args.ids:make_report(root)
        if not args.ids:record('execution_finished',scope+'-complete','所选'+str(len(data))+'题双模型运行完成；成绩见complete.json，待最终结果审阅',str((root/'complete.json').relative_to(ROOT)),0)
        print(json.dumps({'event':'complete',**summary}),flush=True)
    finally:calc.close()
if __name__=='__main__':
    try:main()
    except BaseException as exc:
        if '--out' in sys.argv:
            failed_root=Path(sys.argv[sys.argv.index('--out')+1]).resolve()
            dump(failed_root/'failed.json',{'error':repr(exc),'pid':os.getpid(),'time':time.strftime('%Y-%m-%dT%H:%M:%SZ',time.gmtime())})
            if '--ids' not in sys.argv:
                record('execution_failed','benchmark-run','评测进程已停止：'+repr(exc),str((failed_root/'failed.json').relative_to(ROOT)),1)
        raise
