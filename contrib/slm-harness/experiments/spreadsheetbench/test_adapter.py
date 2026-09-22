import json
from pathlib import Path
import tempfile
import unittest
from adapter import execute_python, resolve_regions, compare_output
import openpyxl

class AdapterTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
        self.input=self.root/'input.xlsx';w=openpyxl.Workbook();w.active['A1']=2;w.save(self.input)
    def tearDown(self):self.temp.cleanup()
    def test_read_only_inspection_cannot_change_work_or_create_output(self):
        work=self.root/'readonly';work.mkdir();(work/'keep.txt').write_text('unchanged')
        for code in ["open('/work/keep.txt','w').write('bad')", "open('/work/output.xlsx','w').write('bad')"]:
            result=execute_python(self.input,work,code,read_only=True)
            self.assertNotEqual(result['exit_code'],0)
        self.assertEqual((work/'keep.txt').read_text(),'unchanged')
        self.assertFalse((work/'output.xlsx').exists())
    def test_isolation(self):
        secret=self.root/'hidden-gold.txt';secret.write_text('secret')
        code=f"import os,socket\nassert not os.path.exists({str(secret)!r})\nassert not os.path.exists('/home/hugo/.ssh')\ns=socket.socket(); s.settimeout(.2)\ntry:\n s.connect(('1.1.1.1',80)); raise AssertionError('network exposed')\nexcept OSError: pass\nprint('isolated')"
        result=execute_python(self.input,self.root/'work',code)
        self.assertEqual(result['exit_code'],0,result);self.assertIn('isolated',result['stdout'])
    def test_fresh_output_and_formula_grade(self):
        work=self.root/'work';code="import openpyxl\nw=openpyxl.load_workbook('/input/input.xlsx');w.active['B1']=4;w.save('/work/output.xlsx')"
        result=execute_python(self.input,work,code, fresh=True)
        self.assertEqual(result['exit_code'],0,result)
        self.assertTrue((work/'output.xlsx').is_file())
        self.assertTrue(compare_output(work/'output.xlsx',work/'output.xlsx',{'answer_position':'B1','instruction_type':'Cell-Level Manipulation'})['passed'])
        self.assertFalse(compare_output(self.input,work/'output.xlsx',{'answer_position':'B1','instruction_type':'Cell-Level Manipulation'})['passed'])
        execute_python(self.input,work,"print('no output')",fresh=True)
        self.assertFalse((work/'output.xlsx').exists())
    def test_regions(self):
        w=openpyxl.Workbook();w.active.title='Other';w.create_sheet('Correct')['B3']=1
        self.assertEqual(resolve_regions({'answer_position':'A:B','answer_sheet':'Correct'},w,w),[('Correct','A1:B3')])
        self.assertEqual(resolve_regions({'answer_position':"Correct'!B3"},w,w),[('Correct','B3:B3')])
    def test_sheet_with_commas(self):
        w=openpyxl.Workbook();w.active.title='b2b, sez, de'
        self.assertEqual(resolve_regions({'answer_position':"'b2b, sez, de'!A5:V10"},w,w),[('b2b, sez, de','A5:V10')])
    def test_host_never_follows_work_symlinks(self):
        victim=self.root/'host-file';victim.write_text('must stay unchanged')
        work=self.root/'work'
        execute_python(self.input,work,f"import os\nfor p in ['.execute.py','.stdout','.stderr']:\n try: os.unlink('/work/'+p)\n except FileNotFoundError: pass\n os.symlink({str(victim)!r},'/work/'+p)")
        execute_python(self.input,work,"print('second inspection')")
        self.assertEqual(victim.read_text(),'must stay unchanged')
    def test_resume_archives_stale_attempt(self):
        from run import prepare_attempt
        task=self.root/'run'/'9b'/'task';task.mkdir(parents=True)
        (task/'solution.py').write_text('stale solution')
        (task/'llm-trace.jsonl').write_text('old trace')
        prepare_attempt(task)
        self.assertEqual(list(task.iterdir()),[])
        archived=list((self.root/'run'/'.interrupted'/'9b').glob('*/solution.py'))
        self.assertEqual(len(archived),1)
        self.assertEqual(archived[0].read_text(),'stale solution')
    def test_reject_symlink(self):
        result=execute_python(self.input,self.root/'work',"import os;os.symlink('/input/input.xlsx','/work/output.xlsx')",fresh=True)
        self.assertEqual(result['exit_code'],0)
        self.assertFalse(result['output_valid'])
if __name__=='__main__':unittest.main()
