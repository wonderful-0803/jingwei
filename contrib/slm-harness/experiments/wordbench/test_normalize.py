import tempfile
import unittest
from pathlib import Path

from docx import Document

from normalize import normalize


class NormalizeTests(unittest.TestCase):
    def test_normalizes_cjk_paragraphs_styles_tables_and_sections(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'sample.docx'
            doc = Document()
            doc.add_paragraph('标题', 'Heading 1')
            doc.add_paragraph('中文段落')
            table = doc.add_table(rows=1, cols=2)
            table.cell(0, 0).text = '键'
            table.cell(0, 1).text = '值'
            doc.save(path)
            actual = normalize(path)
            self.assertEqual(actual['paragraphs'][0]['text'], '标题')
            self.assertEqual(actual['paragraphs'][0]['style'], 'Heading 1')
            self.assertEqual(actual['tables'][0]['rows'], [['键', '值']])
            self.assertEqual(actual['sections'][0]['orientation'], 'portrait')
            self.assertIn('word/document.xml', actual['parts'])


if __name__ == '__main__':
    unittest.main()
