import unittest,tempfile,shutil
from pathlib import Path
from docx import Document
from tasks import build_document,build_task_specs
from grader import grade
from word_plugin import apply_operations
class Boundaries(unittest.TestCase):
 def test_untouched_input_is_not_a_solution_for_any_task(self):
  with tempfile.TemporaryDirectory() as temp:
   for spec in build_task_specs():
    with self.subTest(id=spec['id']):
     inp=Path(temp)/'in.docx'; gold=Path(temp)/'gold.docx'
     build_document(spec,inp);build_document(spec,gold,True)
     r=grade(spec,inp,gold,inp)
     self.assertFalse(r['structure_passed'] and r['content_passed'])
 def test_gold_meets_all_rules(self):
  with tempfile.TemporaryDirectory() as temp:
   for spec in build_task_specs():
    with self.subTest(id=spec['id']):
     inp=Path(temp)/'in.docx'; gold=Path(temp)/'gold.docx'
     build_document(spec,inp);build_document(spec,gold,True)
     r=grade(spec,inp,gold,gold)
     self.assertTrue(r['structure_passed'] and r['content_passed'],r)
 def test_ambiguous_aliases_rejected_and_previous_output_unchanged(self):
  with tempfile.TemporaryDirectory() as temp:
   inp=Path(temp)/'in.docx';out=Path(temp)/'out.docx'
   Document().save(inp);out.write_bytes(b'previous')
   with self.assertRaises(ValueError):apply_operations(inp,out,{'operations':[{'op':'replace_text','text':'new'}]})
   self.assertEqual(out.read_bytes(),b'previous')
 def test_no_placeholder_sources(self):
  with tempfile.TemporaryDirectory() as temp:
   for spec in build_task_specs():
    p=Path(temp)/'in.docx';build_document(spec,p)
    self.assertNotIn('这是用于 Word benchmark 的合成输入。','\n'.join(p.text for p in Document(p).paragraphs),spec['id'])
