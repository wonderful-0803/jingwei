import tempfile
import unittest
from pathlib import Path

from docx import Document

from tasks import build_task_specs, build_document
from word_plugin import apply_operations, snapshot


class WordPluginTests(unittest.TestCase):
    def test_snapshot_and_replace_are_isolated(self):
        spec = next(s for s in build_task_specs() if s['id'] == 'fidelity-protected')
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / 'input.docx'; output = Path(directory) / 'output.docx'
            build_document(spec, source)
            result = apply_operations(source, output, {'operations': [{'op': 'replace_text', 'old': '旧术语', 'new': '新术语'}]})
            self.assertEqual(result['operation_count'], 1)
            self.assertIn('新术语', '\n'.join(p['text'] for p in snapshot(output)['paragraphs']))
            self.assertIn('旧术语', '\n'.join(p['text'] for p in snapshot(source)['paragraphs']))

    def test_table_operations(self):
        spec = next(s for s in build_task_specs() if s['id'] == 'table-fill')
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / 'input.docx'; output = Path(directory) / 'output.docx'
            build_document(spec, source)
            apply_operations(source, output, {'operations': [{'op': 'set_cell', 'row_key': 'Alpha', 'col_index': 1, 'value': '已完成'}]})
            self.assertEqual(Document(output).tables[0].cell(1, 1).text, '已完成')

    def test_model_aliases_are_normalized(self):
        spec = next(s for s in build_task_specs() if s['id'] == 'fidelity-protected')
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / 'input.docx'; output = Path(directory) / 'output.docx'
            build_document(spec, source)
            apply_operations(source, output, {'operations': [{'op': 'replace_text', 'texts': ['术语：旧术语'], 'text': '术语：新术语'}]})
            self.assertIn('新术语', '\n'.join(p.text for p in Document(output).paragraphs))

    def test_style_name_alias_applies_to_all_matching_paragraphs(self):
        spec = next(s for s in build_task_specs() if s['id'] == 'style-unify')
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / 'input.docx'; output = Path(directory) / 'output.docx'
            build_document(spec, source)
            apply_operations(source, output, {'operations': [{'op': 'set_style', 'match': 'Quote', 'style': 'Normal'}]})
            self.assertEqual(Document(output).paragraphs[2].style.name, 'Normal')

    def test_replace_text_accepts_text_as_new_alias(self):
        spec = next(s for s in build_task_specs() if s['id'] == 'delivery-brief')
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / 'input.docx'; output = Path(directory) / 'output.docx'
            build_document(spec, source)
            apply_operations(source, output, {'operations': [{'op': 'replace_text', 'old': '项目资料', 'text': '项目简报'}]})
            self.assertEqual(Document(output).paragraphs[0].text, '项目简报')


if __name__ == '__main__': unittest.main()
