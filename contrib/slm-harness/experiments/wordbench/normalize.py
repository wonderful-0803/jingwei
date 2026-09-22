"""Stable, gold-independent docx representation for grading."""
import hashlib
import json
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT


def _runs(paragraph):
    return [{'text': run.text, 'bold': bool(run.bold), 'italic': bool(run.italic)}
            for run in paragraph.runs]


def _paragraph(paragraph):
    return {'text': paragraph.text, 'style': paragraph.style.name,
            'runs': _runs(paragraph),
            'page_break_before': bool(paragraph.paragraph_format.page_break_before)}


def _section(section):
    orientation = 'landscape' if section.orientation == WD_ORIENT.LANDSCAPE else 'portrait'
    return {'orientation': orientation,
            'page_width': section.page_width.twips,
            'page_height': section.page_height.twips,
            'margins': {name: getattr(section, name).twips for name in
                        ('top_margin', 'bottom_margin', 'left_margin', 'right_margin')}}


def normalize(path):
    path = Path(path)
    doc = Document(path)
    paragraphs = [_paragraph(p) for p in doc.paragraphs]
    tables = [{'rows': [[cell.text for cell in row.cells] for row in table.rows],
               'style': table.style.name if table.style else None}
              for table in doc.tables]
    headers = [[_paragraph(p) for p in section.header.paragraphs] for section in doc.sections]
    footers = [[_paragraph(p) for p in section.footer.paragraphs] for section in doc.sections]
    with zipfile.ZipFile(path) as archive:
        parts = sorted(archive.namelist())
        media = sorted(name for name in parts if name.startswith('word/media/'))
        rels_hash = hashlib.sha256(b''.join(archive.read(name) for name in parts
                                             if name.endswith('.rels'))).hexdigest()
    return {'paragraphs': paragraphs, 'tables': tables, 'sections': [_section(s) for s in doc.sections],
            'headers': headers, 'footers': footers, 'media_count': len(media),
            'media_parts': media, 'parts': parts, 'rels_sha256': rels_hash}


def text_content(normalized):
    values = [p['text'] for p in normalized['paragraphs']]
    values += [cell for table in normalized['tables'] for row in table['rows'] for cell in row]
    return '\n'.join(values)


def dump_normalized(path, destination):
    Path(destination).write_text(json.dumps(normalize(path), ensure_ascii=False, indent=2) + '\n')
